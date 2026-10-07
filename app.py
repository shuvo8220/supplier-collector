"""
Supplier Collector App - MongoDB Version
Flask backend with MongoDB, image upload support, and Excel export
"""
from flask import Flask, jsonify, request, send_file, send_from_directory
from flask_cors import CORS
from datetime import datetime
from bson import ObjectId
import json
import os
from pathlib import Path

# Import custom modules
from db_config import get_db, init_db
from upload_handler import (
    save_uploaded_file, 
    init_upload_folders, 
    get_file_url,
    UPLOAD_CATEGORIES
)

app = Flask(__name__, static_folder='static')
CORS(app)

# Initialize upload folders
init_upload_folders()

# Get MongoDB connection
db = get_db()

def load_schema():
    """Load schema from JSON file"""
    with open('schema.json', 'r', encoding='utf-8') as f:
        return json.load(f)

@app.route('/')
def index():
    """Serve main page"""
    return send_file('static/index.html')

@app.route('/api/schema')
def get_schema():
    """Return schema for frontend"""
    return jsonify(load_schema())

@app.route('/api/upload', methods=['POST'])
def upload_file():
    """
    Handle file uploads
    Expects: file, category, company (optional)
    Returns: {success, filepath, filename, ...}
    """
    try:
        # Check if file is present
        if 'file' not in request.files:
            return jsonify({'success': False, 'error': 'No file provided'}), 400
        
        file = request.files['file']
        category = request.form.get('category', 'document')
        company = request.form.get('company', None)
        
        # Validate category
        if category not in UPLOAD_CATEGORIES:
            return jsonify({
                'success': False, 
                'error': f'Invalid category. Must be one of: {list(UPLOAD_CATEGORIES.keys())}'
            }), 400
        
        # Save file
        result = save_uploaded_file(file, category, company)
        
        if result['success']:
            return jsonify(result), 200
        else:
            return jsonify(result), 400
            
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/uploads/<path:filepath>')
def serve_upload(filepath):
    """Serve uploaded files"""
    return send_from_directory('uploads', filepath)

@app.route('/api/submit/<form_id>', methods=['POST'])
def submit_form(form_id):
    """
    Submit form data to MongoDB
    Creates or updates supplier and related data across multiple collections
    """
    try:
        data = request.json
        company_name = data.get('company', '').strip()
        
        if not company_name:
            return jsonify({'error': 'Company name is required'}), 400
        
        print(f"\n[SUBMIT] Form: {form_id}, Company: {company_name}")
        
        # Check if supplier exists
        supplier = db.suppliers.find_one({'company': company_name})
        
        if supplier:
            supplier_id = supplier['_id']
            print(f"[UPDATE] Existing supplier: {supplier_id}")
        else:
            # Create new supplier
            supplier_doc = {
                'company': company_name,
                'date': data.get('date'),
                'collected_by': data.get('collected_by'),
                'est_year': data.get('est_year'),
                'employees': data.get('employees'),
                'business_type': data.get('business_type'),
                'created_at': datetime.utcnow(),
                'updated_at': datetime.utcnow()
            }
            result = db.suppliers.insert_one(supplier_doc)
            supplier_id = result.inserted_id
            print(f"[CREATE] New supplier: {supplier_id}")
        
        # Update supplier basic info if provided
        if form_id == 'supplier':
            db.suppliers.update_one(
                {'_id': supplier_id},
                {'$set': {
                    'date': data.get('date'),
                    'collected_by': data.get('collected_by'),
                    'est_year': data.get('est_year'),
                    'employees': data.get('employees'),
                    'business_type': data.get('business_type'),
                    'updated_at': datetime.utcnow()
                }}
            )
        
        # Save product data
        product_fields = [
            'product_type', 'product_type_custom', 'main_products', 'main_products_other',
            'material', 'material_source', 'material_quality', 'material_supplier_company',
            'product_working_status', 'product_origin', 'custom', 'custom_details'
        ]
        product_data = {k: data.get(k) for k in product_fields if k in data}
        if product_data:
            product_data['supplier_id'] = supplier_id
            product_data['updated_at'] = datetime.utcnow()
            
            # Upsert (update or insert)
            db.products.update_one(
                {'supplier_id': supplier_id},
                {'$set': product_data, '$setOnInsert': {'created_at': datetime.utcnow()}},
                upsert=True
            )
            print(f"[SAVED] Product data")
        
        # Save sample data
        sample_fields = [
            'sample', 'sample_details', 'sample_type', 'sample_quality',
            'sample_manufacture_date', 'sample_photo_date', 'sample_images'
        ]
        sample_data = {k: data.get(k) for k in sample_fields if k in data}
        if sample_data:
            sample_data['supplier_id'] = supplier_id
            sample_data['updated_at'] = datetime.utcnow()
            
            db.samples.update_one(
                {'supplier_id': supplier_id},
                {'$set': sample_data, '$setOnInsert': {'created_at': datetime.utcnow()}},
                upsert=True
            )
            print(f"[SAVED] Sample data")
        
        # Save document/verification data
        document_fields = [
            'tin_provided', 'tin_bin', 'trade_license', 'trade_license_no', 'trade_license_image',
            'bsti', 'bsti_image', 'export_exp', 'company_authenticity_verified', 'verification_notes'
        ]
        document_data = {k: data.get(k) for k in document_fields if k in data}
        if document_data:
            document_data['supplier_id'] = supplier_id
            document_data['updated_at'] = datetime.utcnow()
            
            db.documents.update_one(
                {'supplier_id': supplier_id},
                {'$set': document_data, '$setOnInsert': {'created_at': datetime.utcnow()}},
                upsert=True
            )
            print(f"[SAVED] Document data")
        
        # Save commercial data
        commercial_fields = [
            'opening_time', 'closing_time', 'moq', 'capacity', 'capacity_per_order_max',
            'capacity_per_load', 'lead_time', 'delivery_cost', 'vehicle_available',
            'price_list', 'price_list_details', 'price_list_images', 'price_range'
        ]
        commercial_data = {k: data.get(k) for k in commercial_fields if k in data}
        if commercial_data:
            commercial_data['supplier_id'] = supplier_id
            commercial_data['updated_at'] = datetime.utcnow()
            
            db.commercial.update_one(
                {'supplier_id': supplier_id},
                {'$set': commercial_data, '$setOnInsert': {'created_at': datetime.utcnow()}},
                upsert=True
            )
            print(f"[SAVED] Commercial data")
        
        # Save contact data
        contact_fields = [
            'contact_person_1_name', 'contact_person_1_designation', 'contact_person_1_phone', 'contact_person_1_whatsapp',
            'contact_person_2_name', 'contact_person_2_designation', 'contact_person_2_phone', 'contact_person_2_whatsapp',
            'contact_person_3_name', 'contact_person_3_designation', 'contact_person_3_phone', 'contact_person_3_whatsapp',
            'email', 'facebook', 'website',
            'factory_addr', 'factory_map_link', 'office_addr', 'office_map_link',
            'factory_own', 'delivery'
        ]
        contact_data = {k: data.get(k) for k in contact_fields if k in data}
        if contact_data:
            contact_data['supplier_id'] = supplier_id
            contact_data['updated_at'] = datetime.utcnow()
            
            db.contacts.update_one(
                {'supplier_id': supplier_id},
                {'$set': contact_data, '$setOnInsert': {'created_at': datetime.utcnow()}},
                upsert=True
            )
            print(f"[SAVED] Contact data")
        
        # Save media data
        media_fields = [
            'photo_product', 'product_images', 'photo_factory', 'factory_images',
            'catalog', 'catalog_images', 'consent', 'signature', 'sign_date'
        ]
        media_data = {k: data.get(k) for k in media_fields if k in data}
        if media_data:
            media_data['supplier_id'] = supplier_id
            media_data['updated_at'] = datetime.utcnow()
            
            db.media.update_one(
                {'supplier_id': supplier_id},
                {'$set': media_data, '$setOnInsert': {'created_at': datetime.utcnow()}},
                upsert=True
            )
            print(f"[SAVED] Media data")
        
        # Save interview data (27 questions)
        interview_fields = [f'q{i}' for i in range(1, 28)] + ['interest', 'next_action', 'notes']
        interview_data = {k: data.get(k) for k in interview_fields if k in data}
        if interview_data:
            interview_data['supplier_id'] = supplier_id
            interview_data['updated_at'] = datetime.utcnow()
            
            db.interviews.update_one(
                {'supplier_id': supplier_id},
                {'$set': interview_data, '$setOnInsert': {'created_at': datetime.utcnow()}},
                upsert=True
            )
            print(f"[SAVED] Interview data")
        
        return jsonify({
            'success': True,
            'message': 'Data saved successfully',
            'supplier_id': str(supplier_id),
            'updated': supplier is not None
        }), 200
        
    except Exception as e:
        print(f"[ERROR] {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

@app.route('/api/suppliers', methods=['GET'])
def get_suppliers():
    """Get all suppliers list"""
    try:
        suppliers = list(db.suppliers.find({}, {
            'company': 1,
            'date': 1,
            'collected_by': 1,
            'business_type': 1,
            'created_at': 1
        }).sort('created_at', -1))
        
        # Convert ObjectId to string
        for s in suppliers:
            s['_id'] = str(s['_id'])
        
        return jsonify({'suppliers': suppliers}), 200
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/supplier/<supplier_id>', methods=['GET'])
def get_supplier_detail(supplier_id):
    """Get complete supplier details with all related data"""
    try:
        # Get supplier
        supplier = db.suppliers.find_one({'_id': ObjectId(supplier_id)})
        if not supplier:
            return jsonify({'error': 'Supplier not found'}), 404
        
        # Convert ObjectId to string
        supplier['_id'] = str(supplier['_id'])
        
        # Get related data
        product = db.products.find_one({'supplier_id': ObjectId(supplier_id)})
        sample = db.samples.find_one({'supplier_id': ObjectId(supplier_id)})
        document = db.documents.find_one({'supplier_id': ObjectId(supplier_id)})
        commercial = db.commercial.find_one({'supplier_id': ObjectId(supplier_id)})
        contact = db.contacts.find_one({'supplier_id': ObjectId(supplier_id)})
        media = db.media.find_one({'supplier_id': ObjectId(supplier_id)})
        interview = db.interviews.find_one({'supplier_id': ObjectId(supplier_id)})
        
        # Convert ObjectIds in related data
        for doc in [product, sample, document, commercial, contact, media, interview]:
            if doc:
                doc['_id'] = str(doc['_id'])
                doc['supplier_id'] = str(doc['supplier_id'])
        
        return jsonify({
            'supplier': supplier,
            'product': product,
            'sample': sample,
            'document': document,
            'commercial': commercial,
            'contact': contact,
            'media': media,
            'interview': interview
        }), 200
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/download', methods=['GET'])
def download_excel():
    """Export MongoDB data to Excel file"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Suppliers"
        
        # Get all suppliers with related data
        suppliers = list(db.suppliers.find().sort('created_at', -1))
        
        if not suppliers:
            return jsonify({'error': 'No data to export'}), 404
        
        # Define columns
        columns = [
            'Company Name', 'Date', 'Collected By', 'Est Year', 'Employees', 'Business Type',
            # Products
            'Product Type', 'Main Products', 'Material', 'Material Source', 'Material Quality',
            'Material Supplier', 'Product Status', 'Product Origin', 'Custom Available', 'Custom Details',
            # Sample
            'Sample Available', 'Sample Type', 'Sample Quality', 'Sample Details',
            # Commercial
            'Opening Time', 'Closing Time', 'MOQ', 'Monthly Capacity', 'Max Order Capacity',
            'Capacity Per Load', 'Lead Time', 'Delivery Cost', 'Vehicle Available',
            'Price List Available', 'Price Range',
            # Contact
            'Email', 'Phone', 'WhatsApp', 'Facebook', 'Website',
            'Factory Address', 'Factory Map Link', 'Office Address', 'Office Map Link',
            'Factory Own', 'Delivery Area',
            # Documents
            'TIN Provided', 'TIN/BIN', 'Trade License', 'Trade License No', 'BSTI',
            'Export Experience', 'Company Verified',
            # Interview (selected questions)
            'Main Buyers', 'Pain Points', 'B2B Marketplace Used',
            # Internal
            'Interest Level', 'Notes', 'Created At'
        ]
        
        # Header row
        for col_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(1, col_idx, col_name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="0F4C81", end_color="0F4C81", fill_type="solid")
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        # Data rows
        for row_idx, supplier in enumerate(suppliers, start=2):
            supplier_id = supplier['_id']
            
            # Get related data
            product = db.products.find_one({'supplier_id': supplier_id}) or {}
            sample = db.samples.find_one({'supplier_id': supplier_id}) or {}
            commercial = db.commercial.find_one({'supplier_id': supplier_id}) or {}
            contact = db.contacts.find_one({'supplier_id': supplier_id}) or {}
            document = db.documents.find_one({'supplier_id': supplier_id}) or {}
            interview = db.interviews.find_one({'supplier_id': supplier_id}) or {}
            
            # Combine data
            row_data = [
                supplier.get('company', ''),
                supplier.get('date', ''),
                supplier.get('collected_by', ''),
                supplier.get('est_year', ''),
                supplier.get('employees', ''),
                supplier.get('business_type', ''),
                # Products
                product.get('product_type', ''),
                ', '.join(product.get('main_products', [])) if isinstance(product.get('main_products'), list) else product.get('main_products', ''),
                product.get('material', ''),
                product.get('material_source', ''),
                product.get('material_quality', ''),
                product.get('material_supplier_company', ''),
                product.get('product_working_status', ''),
                product.get('product_origin', ''),
                product.get('custom', ''),
                product.get('custom_details', ''),
                # Sample
                sample.get('sample', ''),
                sample.get('sample_type', ''),
                sample.get('sample_quality', ''),
                sample.get('sample_details', ''),
                # Commercial
                commercial.get('opening_time', ''),
                commercial.get('closing_time', ''),
                commercial.get('moq', ''),
                commercial.get('capacity', ''),
                commercial.get('capacity_per_order_max', ''),
                commercial.get('capacity_per_load', ''),
                commercial.get('lead_time', ''),
                commercial.get('delivery_cost', ''),
                commercial.get('vehicle_available', ''),
                commercial.get('price_list', ''),
                commercial.get('price_range', ''),
                # Contact
                contact.get('email', ''),
                contact.get('contact_person_1_phone', ''),
                contact.get('contact_person_1_whatsapp', ''),
                contact.get('facebook', ''),
                contact.get('website', ''),
                contact.get('factory_addr', ''),
                contact.get('factory_map_link', ''),
                contact.get('office_addr', ''),
                contact.get('office_map_link', ''),
                contact.get('factory_own', ''),
                ', '.join(contact.get('delivery', [])) if isinstance(contact.get('delivery'), list) else contact.get('delivery', ''),
                # Documents
                document.get('tin_provided', ''),
                document.get('tin_bin', ''),
                document.get('trade_license', ''),
                document.get('trade_license_no', ''),
                document.get('bsti', ''),
                document.get('export_exp', ''),
                document.get('company_authenticity_verified', ''),
                # Interview (selected)
                interview.get('q4', ''),
                interview.get('q9', ''),
                interview.get('q25', ''),
                # Internal
                interview.get('interest', ''),
                interview.get('notes', ''),
                supplier.get('created_at', '').strftime('%Y-%m-%d %H:%M') if supplier.get('created_at') else ''
            ]
            
            for col_idx, value in enumerate(row_data, start=1):
                ws.cell(row_idx, col_idx, value)
        
        # Auto-adjust column widths
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width
        
        # Save to file
        excel_path = 'data/export.xlsx'
        Path('data').mkdir(exist_ok=True)
        wb.save(excel_path)
        
        return send_file(
            excel_path,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'suppliers_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        )
        
    except Exception as e:
        print(f"[ERROR] Excel export: {str(e)}")
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500

if __name__ == '__main__':
    print("\n" + "="*50)
    print("🚀 Supplier Collector App - MongoDB Version")
    print("="*50)
    print("✓ MongoDB connected")
    print("✓ Upload folders ready")
    print("✓ Server starting on http://localhost:5000")
    print("="*50 + "\n")
    
    app.run(debug=True, host='0.0.0.0', port=5000)
